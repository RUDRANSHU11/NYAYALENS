"""Upload validation and text extraction, including the checks that exist for
security reasons (magic bytes, size, zip expansion)."""

import io
import zipfile

import pytest

from backend import config
from backend.document_processing.extract import DocumentError, extract


def build_pdf(text: str, pages: int = 1) -> bytes:
    import fitz

    document = fitz.open()
    for _ in range(pages):
        page = document.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 750), text, fontsize=11)
    return document.tobytes()


def build_docx(paragraphs: list[str], heading: str | None = None) -> bytes:
    from docx import Document

    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_txt_extraction_drops_blank_lines(contract_v1):
    extracted = extract("contract.txt", contract_v1)
    assert extracted.pages == 0
    assert all(line.text.strip() for line in extracted.lines)
    assert any("30 days written notice" in line.text for line in extracted.lines)


def test_pdf_extraction_records_page_numbers():
    data = build_pdf("1. NOTICE PERIOD\nEither party may give 30 days written notice.", pages=2)
    extracted = extract("contract.pdf", data)
    assert extracted.pages == 2
    assert {line.page for line in extracted.lines} == {1, 2}
    assert not extracted.needs_ocr


def test_pdf_without_a_text_layer_is_flagged_for_ocr():
    data = build_pdf("", pages=1)
    extracted = extract("scan.pdf", data)
    assert extracted.needs_ocr
    assert extracted.page_images and extracted.page_images[0][:4] == b"\x89PNG"


def test_docx_extraction_marks_heading_styles():
    data = build_docx(["1. RENT", "The Tenant shall pay Rs. 25,000 per month."], heading="LEASE")
    extracted = extract("lease.docx", data)
    assert extracted.lines[0].heading is True
    assert any("Rs. 25,000" in line.text for line in extracted.lines)


def test_docx_tables_are_extracted():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Security deposit"
    table.rows[0].cells[1].text = "Rs. 50,000"
    buffer = io.BytesIO()
    document.save(buffer)

    extracted = extract("lease.docx", buffer.getvalue())
    assert any("Security deposit | Rs. 50,000" == line.text for line in extracted.lines)


def test_unsupported_extension_is_refused():
    with pytest.raises(DocumentError) as error:
        extract("payload.exe", b"MZ\x90\x00")
    assert error.value.status_code == 415


def test_extension_lies_about_content(tmp_path):
    with pytest.raises(DocumentError) as error:
        extract("contract.pdf", b"<html>not a pdf</html>")
    assert error.value.status_code == 415


def test_empty_upload_is_refused():
    with pytest.raises(DocumentError):
        extract("contract.txt", b"")


def test_oversize_upload_is_refused(monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 32)
    with pytest.raises(DocumentError) as error:
        extract("contract.txt", b"x" * 64)
    assert error.value.status_code == 413


def test_too_many_pages_is_refused(monkeypatch):
    monkeypatch.setattr(config, "MAX_PAGES", 1)
    with pytest.raises(DocumentError) as error:
        extract("contract.pdf", build_pdf("clause text", pages=2))
    assert error.value.status_code == 413


def test_zip_that_is_not_a_docx_is_refused():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a word document")
    with pytest.raises(DocumentError) as error:
        extract("contract.docx", buffer.getvalue())
    assert error.value.status_code == 415


def test_zip_bomb_is_refused(monkeypatch):
    from backend.document_processing import extract as extract_module

    monkeypatch.setattr(extract_module, "_MAX_DOCX_UNCOMPRESSED", 1024)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 100_000)
    with pytest.raises(DocumentError) as error:
        extract("bomb.docx", buffer.getvalue())
    assert error.value.status_code == 413
